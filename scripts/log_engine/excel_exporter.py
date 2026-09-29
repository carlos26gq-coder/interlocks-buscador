"""
SOLVI - Log Engine: Excel Multi-Sheet Audit Report Generator
Exports Linac forensics and telemetry data into professional multi-sheet .xlsx workbooks.
"""

import io
import datetime
from typing import Dict, Any, Union, Optional

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def _apply_header_style(cell, text: str):
    """Applies Elekta Medical Navy Blue header styling."""
    cell.value = text
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _apply_cell_border(cell):
    """Applies a clean thin border to cells."""
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    cell.border = thin_border


def _autofit_columns(ws, max_len_cap: int = 60):
    """Automatically adjusts column widths based on cell content length."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), max_len_cap)


class LinacExcelExporter:
    """Generates multi-sheet Excel workbooks for clinical medical physics and biomedical engineering audits."""

    @staticmethod
    def is_available() -> bool:
        return OPENPYXL_AVAILABLE

    @staticmethod
    def export(analysis_data: Dict[str, Any], output: Optional[Union[str, io.BytesIO]] = None) -> Union[str, io.BytesIO]:
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl no está instalado en el entorno. No se puede generar el archivo .xlsx.")

        wb = openpyxl.Workbook()
        # Remove default sheet
        default_sheet = wb.active
        
        # 1. Executive Summary Sheet
        ws_summary = wb.create_sheet(title="Resumen_Ejecutivo")
        LinacExcelExporter._build_summary_sheet(ws_summary, analysis_data)

        # 2. Hardware Interlocks Sheet
        ws_interlocks = wb.create_sheet(title="Interlocks_Hardware")
        LinacExcelExporter._build_interlocks_sheet(ws_interlocks, analysis_data.get("interlocks", {}))

        # 3. MLC Collisions Sheet
        ws_collisions = wb.create_sheet(title="Colisiones_MLC_Agility")
        LinacExcelExporter._build_collisions_sheet(ws_collisions, analysis_data.get("interlocks", {}))

        # 4. Treatment Deliveries (TRF)
        ws_trf = wb.create_sheet(title="Tratamientos_TRF")
        LinacExcelExporter._build_trf_sheet(ws_trf, analysis_data.get("treatments", {}))

        # 5. Clinical Audit Trail
        ws_audit = wb.create_sheet(title="Auditoria_Clinica")
        LinacExcelExporter._build_audit_sheet(ws_audit, analysis_data.get("audit_trail", {}))

        # 6. CCP Supervisor Logs
        ws_ccp = wb.create_sheet(title="Supervisor_CCP")
        LinacExcelExporter._build_ccp_sheet(ws_ccp, analysis_data.get("supervisor", {}))

        # 7. Optical Calibration
        ws_opt = wb.create_sheet(title="Calibracion_Optica")
        LinacExcelExporter._build_optical_sheet(ws_opt, analysis_data.get("optical", {}))

        # Remove initial empty sheet
        if default_sheet in wb.worksheets:
            wb.remove(default_sheet)

        # Save to file path or in-memory BytesIO
        if output is None:
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return buf
        elif isinstance(output, str):
            wb.save(output)
            return output
        else:
            wb.save(output)
            output.seek(0)
            return output

    @staticmethod
    def _build_summary_sheet(ws, data: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        exec_s = data.get("executive_summary", {})
        profile = data.get("profile", {})
        correlations = data.get("correlations", [])

        # Title Banner
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "SOLVI — INFORME DE AUDITORÍA FORENSE Y DIAGNÓSTICO DE ACELERADOR LINEAL ELEKTA"
        title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # Subtitle
        ws["A2"].value = f"Fecha de emisión: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Análisis automatizado de integridad de logs"
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="595959")
        ws.row_dimensions[2].height = 20

        # Section 1: Linac Profile
        ws["A4"].value = "1. IDENTIFICACIÓN Y ESPECIFICACIONES DEL LINAC"
        ws["A4"].font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

        hw_opts = ", ".join(profile.get("hardware_options", ["Agility 160 MLC", "Cuña Motorizada", "Servo Cañón Avanzado"]))
        profile_rows = [
            ("Número de Serie (Linac ID):", exec_s.get("linac_id", "4574"), "Nombre en Red:", exec_s.get("linac_name", "05Elekta")),
            ("Equipo Consola:", exec_s.get("console_host", "ELEKTA5"), "Software de Control:", exec_s.get("software_version", "Integrity 4.0.6")),
            ("Horas de Radiación (HT):", f"{exec_s.get('ht_hours', 0)} h", "Horas Electrónica (LT):", f"{exec_s.get('lt_hours', 0)} h"),
            ("Colimador Multiláminas:", "Elekta Agility (160 Láminas, 80 Pares)", "Escala Coordenadas:", exec_s.get("scale", "IEC1217")),
            ("Energías Autorizadas:", profile.get("energies", "6MV"), "Subsistemas Físicos:", hw_opts),
        ]

        row_idx = 5
        for r in profile_rows:
            ws.cell(row=row_idx, column=1, value=r[0]).font = Font(bold=True, size=10)
            ws.cell(row=row_idx, column=2, value=r[1]).font = Font(size=10, color="1F4E79")
            ws.cell(row=row_idx, column=4, value=r[2]).font = Font(bold=True, size=10)
            ws.cell(row=row_idx, column=5, value=r[3]).font = Font(size=10, color="1F4E79")
            row_idx += 1

        # Section 2: KPIs
        row_idx += 1
        ws.cell(row=row_idx, column=1, value="2. MÉTRICAS GLOBALES DE DIAGNÓSTICO (KPIS DE SALUD)").font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        row_idx += 1

        kpi_headers = ["Métrica Forense", "Valor Registrado", "Interpretación Clínica / Física"]
        for c_idx, h in enumerate(kpi_headers, start=1):
            cell = ws.cell(row=row_idx, column=c_idx)
            _apply_header_style(cell, h)
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        kpis = [
            ("Archivos Totales Analizados", f"{exec_s.get('total_files_scanned', 0):,} archivos ({exec_s.get('total_volume_mb', 0):,} MB)", "Volumen total de telemetría y diagnósticos procesados"),
            ("Haces Entregados (.trf)", f"{exec_s.get('total_beams_delivered', 0):,} tratamientos", "Sesiones y campos de irradiación completados o interrumpidos"),
            ("Dosis Total Administrada", f"{exec_s.get('total_mu_delivered', 0):,.1f} MU", "Unidades Monitor totales emitidas por el Linac"),
            ("Interlocks de Suspensión Hardware", f"{exec_s.get('total_suspension_interlocks', 0):,} eventos", "Eventos donde el hardware del Linac detuvo la radiación por seguridad"),
            ("Avisos de Proximidad / Colisión MLC", f"{exec_s.get('total_mlc_collision_warnings', 0):,} alertas", "Aproximaciones críticas entre láminas opuestas de Agility"),
            ("Reinicios de Máquina (Linac Reset)", f"{exec_s.get('total_linac_resets', 0):,} ejecuciones", "Rearmes manuales de interlocks ejecutados por tecnólogos u operadores"),
            ("Pérdidas de Latido del Controlador", f"{exec_s.get('total_heartbeat_misses', 0):,} caídas", "Retardos o pérdidas de comunicación entre consola y gantry"),
        ]

        for metric, val, desc in kpis:
            c1 = ws.cell(row=row_idx, column=1, value=metric)
            c2 = ws.cell(row=row_idx, column=2, value=val)
            c3 = ws.cell(row=row_idx, column=3, value=desc)
            c1.font = Font(bold=True, size=10)
            c2.font = Font(size=10, bold=True, color="1F4E79")
            c3.font = Font(size=10, italic=True)
            for c in [c1, c2, c3]:
                _apply_cell_border(c)
            row_idx += 1

        # Section 3: Correlations & Root Causes
        row_idx += 1
        ws.cell(row=row_idx, column=1, value="3. ANÁLISIS DE CAUSA RAÍZ Y RECOMENDACIONES DE INGENIERÍA").font = Font(name="Calibri", size=11, bold=True, color="1F4E79")
        row_idx += 1

        corr_headers = ["Subsistema", "Severidad", "Hallazgo Forense", "Impacto Operativo", "Acción Técnica Recomendada"]
        for c_idx, h in enumerate(corr_headers, start=1):
            cell = ws.cell(row=row_idx, column=c_idx)
            _apply_header_style(cell, h)
        ws.row_dimensions[row_idx].height = 24
        row_idx += 1

        for c in correlations:
            c1 = ws.cell(row=row_idx, column=1, value=c.get("subsystem", ""))
            c2 = ws.cell(row=row_idx, column=2, value=c.get("severity", ""))
            c3 = ws.cell(row=row_idx, column=3, value=c.get("finding", ""))
            c4 = ws.cell(row=row_idx, column=4, value=c.get("impact", ""))
            c5 = ws.cell(row=row_idx, column=5, value=c.get("recommendation", ""))
            
            c1.font = Font(bold=True, size=10)
            c2.font = Font(bold=True, size=10, color="C00000" if c.get("severity") == "CRITICAL" else "B25900")
            c3.font = Font(size=9.5)
            c4.font = Font(size=9.5)
            c5.font = Font(size=9.5, bold=True, color="1F4E79")

            for cell in [c1, c2, c3, c4, c5]:
                _apply_cell_border(cell)
            row_idx += 1

        _autofit_columns(ws, max_len_cap=70)

    @staticmethod
    def _build_interlocks_sheet(ws, interlocks: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Fecha / Hora", "Severidad", "Item ID", "Código Interlock", "Valor Crudo", "Componente / Descripción Biomédica"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        events = interlocks.get("events_suspensions", [])
        for r_idx, ev in enumerate(events, start=2):
            c_time = ws.cell(row=r_idx, column=1, value=ev.get("timestamp", ""))
            c_sev = ws.cell(row=r_idx, column=2, value=ev.get("severity", "CRITICAL"))
            c_item = ws.cell(row=r_idx, column=3, value=ev.get("item_id", 0))
            c_code = ws.cell(row=r_idx, column=4, value=ev.get("code", 0))
            c_val = ws.cell(row=r_idx, column=5, value=ev.get("value", 0))
            c_desc = ws.cell(row=r_idx, column=6, value=ev.get("description", ""))

            c_sev.font = Font(bold=True, color="C00000")
            c_item.font = Font(bold=True)
            c_code.font = Font(name="Calibri")
            
            for c in [c_time, c_sev, c_item, c_code, c_val, c_desc]:
                _apply_cell_border(c)

        ws.auto_filter.ref = f"A1:F{max(len(events) + 1, 2)}"
        _autofit_columns(ws)

    @staticmethod
    def _build_collisions_sheet(ws, interlocks: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Fecha / Hora", "Severidad", "Par de Lámina", "Tipo Alerta", "Distancia Mínima (mm)", "Setpoint Modificado (Anti-colisión)", "Diagnóstico"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        events = interlocks.get("events_collisions", [])
        for r_idx, ev in enumerate(events, start=2):
            c_time = ws.cell(row=r_idx, column=1, value=ev.get("timestamp", ""))
            c_sev = ws.cell(row=r_idx, column=2, value=ev.get("severity", "WARNING"))
            c_pair = ws.cell(row=r_idx, column=3, value=f"Par #{ev.get('leaf_pair', 0)}")
            c_type = ws.cell(row=r_idx, column=4, value=f"Tipo {ev.get('collision_type', 0)}")
            c_dist = ws.cell(row=r_idx, column=5, value=ev.get("distance_mm", 0.0))
            c_mod = ws.cell(row=r_idx, column=6, value="SÍ (Activo)" if ev.get("setpoint_modified") else "NO")
            c_desc = ws.cell(row=r_idx, column=7, value=ev.get("description", ""))

            c_sev.font = Font(bold=True, color="B25900")
            c_dist.number_format = "0.000"
            c_pair.font = Font(bold=True)

            for c in [c_time, c_sev, c_pair, c_type, c_dist, c_mod, c_desc]:
                _apply_cell_border(c)

        ws.auto_filter.ref = f"A1:G{max(len(events) + 1, 2)}"
        _autofit_columns(ws)

    @staticmethod
    def _build_trf_sheet(ws, treatments: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Fecha / Hora UTC", "Nombre del Haz / Campo", "Etiqueta", "Dosis (MU)", "Duración (s)", "Canales Registrados", "Estado de Entrega", "Archivo TRF"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        deliveries = treatments.get("deliveries", [])
        for r_idx, d in enumerate(deliveries, start=2):
            c_time = ws.cell(row=r_idx, column=1, value=d.get("date_utc", ""))
            c_name = ws.cell(row=r_idx, column=2, value=d.get("beam_name", ""))
            c_lbl = ws.cell(row=r_idx, column=3, value=d.get("field_label", ""))
            c_mu = ws.cell(row=r_idx, column=4, value=d.get("mu", 0.0))
            c_dur = ws.cell(row=r_idx, column=5, value=d.get("duration_sec", 0.0))
            c_ch = ws.cell(row=r_idx, column=6, value=d.get("channels_count", 0))
            c_st = ws.cell(row=r_idx, column=7, value=d.get("status", ""))
            c_file = ws.cell(row=r_idx, column=8, value=d.get("file_name", ""))

            c_mu.number_format = "0.00"
            c_dur.number_format = "0.00"
            c_st.font = Font(bold=True, color="375623" if d.get("status") == "Completed" else "C00000")

            for c in [c_time, c_name, c_lbl, c_mu, c_dur, c_ch, c_st, c_file]:
                _apply_cell_border(c)

        ws.auto_filter.ref = f"A1:H{max(len(deliveries) + 1, 2)}"
        _autofit_columns(ws)

    @staticmethod
    def _build_audit_sheet(ws, audit: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Fecha / Hora", "Host", "ID Evento", "Nombre de Evento Clínico", "ID Usuario", "Usuario", "ID Linac"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        events = audit.get("events", [])
        for r_idx, ev in enumerate(events, start=2):
            c_time = ws.cell(row=r_idx, column=1, value=ev.get("occurred", ""))
            c_host = ws.cell(row=r_idx, column=2, value=ev.get("host", ""))
            c_eid = ws.cell(row=r_idx, column=3, value=ev.get("event_id", ""))
            c_ename = ws.cell(row=r_idx, column=4, value=ev.get("event_name", ""))
            c_uid = ws.cell(row=r_idx, column=5, value=ev.get("user_id", ""))
            c_uname = ws.cell(row=r_idx, column=6, value=ev.get("user_name", ""))
            c_lid = ws.cell(row=r_idx, column=7, value=ev.get("linac_id", ""))

            if "Reset" in ev.get("event_name", ""):
                c_ename.font = Font(bold=True, color="C00000")
            elif "Logon" in ev.get("event_name", ""):
                c_ename.font = Font(color="1F4E79")

            for c in [c_time, c_host, c_eid, c_ename, c_uid, c_uname, c_lid]:
                _apply_cell_border(c)

        ws.auto_filter.ref = f"A1:G{max(len(events) + 1, 2)}"
        _autofit_columns(ws)

    @staticmethod
    def _build_ccp_sheet(ws, ccp: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Fecha / Hora", "Nivel", "Componente Supervisor", "Mensaje de Diagnóstico"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        events = ccp.get("sample_events", [])
        for r_idx, ev in enumerate(events, start=2):
            c_time = ws.cell(row=r_idx, column=1, value=ev.get("timestamp", ""))
            c_lvl = ws.cell(row=r_idx, column=2, value=ev.get("level", ""))
            c_comp = ws.cell(row=r_idx, column=3, value=ev.get("component", ""))
            c_msg = ws.cell(row=r_idx, column=4, value=ev.get("message", ""))

            lvl_str = ev.get("level", "")
            if lvl_str in ["ERROR", "FATAL"]:
                c_lvl.font = Font(bold=True, color="C00000")
            elif lvl_str == "WARN":
                c_lvl.font = Font(bold=True, color="B25900")

            for c in [c_time, c_lvl, c_comp, c_msg]:
                _apply_cell_border(c)

        ws.auto_filter.ref = f"A1:D{max(len(events) + 1, 2)}"
        _autofit_columns(ws)

    @staticmethod
    def _build_optical_sheet(ws, opt: Dict[str, Any]):
        ws.views.sheetView[0].showGridLines = True
        headers = ["Banco Láminas", "Lámina", "Puntos de Calibración", "Pico Distorsión Óptica (µm)", "Distorsión Media (µm)", "Estado"]
        for c_idx, h in enumerate(headers, start=1):
            _apply_header_style(ws.cell(row=1, column=c_idx), h)
        ws.row_dimensions[1].height = 26

        row_idx = 2
        for bank_data in opt.get("banks", []):
            bank_name = bank_data.get("summary", {}).get("leaf_bank", "Y1")
            for leaf in bank_data.get("leaves", []):
                c_bank = ws.cell(row=row_idx, column=1, value=bank_name)
                c_title = ws.cell(row=row_idx, column=2, value=leaf.get("title", ""))
                c_pts = ws.cell(row=row_idx, column=3, value=leaf.get("calibration_points_count", 0))
                c_peak = ws.cell(row=row_idx, column=4, value=leaf.get("peak_distortion_microns", 0))
                c_mean = ws.cell(row=row_idx, column=5, value=leaf.get("mean_distortion_microns", 0))
                c_st = ws.cell(row=row_idx, column=6, value="Calibrado OK" if leaf.get("peak_distortion_microns", 0) < 250 else "Revisar")

                c_peak.number_format = "#,##0"
                c_mean.number_format = "#,##0.0"
                if leaf.get("peak_distortion_microns", 0) >= 200:
                    c_peak.font = Font(bold=True, color="B25900")

                for c in [c_bank, c_title, c_pts, c_peak, c_mean, c_st]:
                    _apply_cell_border(c)
                row_idx += 1

        ws.auto_filter.ref = f"A1:F{max(row_idx - 1, 2)}"
        _autofit_columns(ws)
